from types import SimpleNamespace
from unittest.mock import AsyncMock
import io
import asyncio
from functools import wraps
import pytest
from utils.attachment_reader import collect_attachments, build_attachment_parts, extract_document


def run_async(test):
    @wraps(test)
    def run(*args, **kwargs):
        return asyncio.run(test(*args, **kwargs))
    return run


@run_async
async def test_reply_fetches_attachment_and_preserves_question():
    att = SimpleNamespace(filename='note.txt', content_type='text/plain', size=5,
                          read=AsyncMock(return_value=b'hello'))
    target = SimpleNamespace(attachments=[att])
    channel = SimpleNamespace(id=10, fetch_message=AsyncMock(return_value=target))
    message = SimpleNamespace(attachments=[], channel=channel,
        reference=SimpleNamespace(channel_id=10, message_id=20))
    attachments = await collect_attachments(message)
    parts = await build_attachment_parts('請摘要', attachments)
    assert parts[0]['text'] == '請摘要'
    assert 'note.txt' in parts[1]['text'] and 'hello' in parts[1]['text']
    assert '內容已截斷' not in parts[1]['text']
    channel.fetch_message.assert_awaited_once_with(20)


@run_async
async def test_cross_channel_reply_rejected():
    message = SimpleNamespace(attachments=[], channel=SimpleNamespace(id=10),
        reference=SimpleNamespace(channel_id=99, message_id=20))
    with pytest.raises(ValueError, match='同一頻道'):
        await collect_attachments(message)


@run_async
async def test_oversized_file_not_downloaded():
    att = SimpleNamespace(filename='big.txt', size=25*1024*1024, read=AsyncMock())
    parts = await build_attachment_parts('摘要', [att])
    assert '超過' in parts[1]['text']
    att.read.assert_not_awaited()


def test_office_extraction():
    from docx import Document
    from openpyxl import Workbook
    doc = Document(); doc.add_paragraph('文件測試')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格內容'
    data = io.BytesIO(); doc.save(data)
    assert '表格內容' in extract_document('a.docx', data.getvalue())
    book = Workbook(); book.active.append(['姓名', '分數']); book.active.append(['小明', 90])
    data = io.BytesIO(); book.save(data)
    assert '小明' in extract_document('a.xlsx', data.getvalue())


def test_pdf_without_text_is_explicit():
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    data = io.BytesIO(); writer.write(data)
    with pytest.raises(ValueError, match='文字'):
        extract_document('a.pdf', data.getvalue())


def test_big5_text_is_not_misread_as_utf16():
    assert extract_document('a.txt', '繁體中文'.encode('big5')) == '繁體中文'


def test_pdf_preserves_truncation_sentinel(monkeypatch):
    import pypdf
    from utils.attachment_reader import MAX_TEXT
    page = SimpleNamespace(extract_text=lambda: '字' * (MAX_TEXT + 100))
    monkeypatch.setattr(pypdf, 'PdfReader', lambda _: SimpleNamespace(is_encrypted=False, pages=[page]))
    assert len(extract_document('a.pdf', b'placeholder')) > MAX_TEXT
